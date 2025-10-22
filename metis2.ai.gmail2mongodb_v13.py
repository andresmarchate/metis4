import base64
import mimetypes
import os
import pickle
from PyPDF2 import PdfReader
from PyPDF2.errors import FileNotDecryptedError
import pytesseract
from PIL import Image
from docx import Document as DocxDocument
import openpyxl
import json
from pymongo import MongoClient, TEXT, ASCENDING
import hashlib
import re
import requests
from collections import OrderedDict, Counter
from sentence_transformers import SentenceTransformer
import numpy as np
import zlib
from googleapiclient.discovery import build
from google_auth_oauthlib.flow import InstalledAppFlow

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']
CREDS_FILE = 'credentials.json'
TOKEN_FILE = 'token.pickle'

# Configuración de Tesseract
pytesseract.pytesseract.tesseract_cmd = r'/usr/bin/tesseract'  # Ajusta según tu sistema

# Archivo y configuración del caché
cache_file = 'mistral_cache.pkl'
cache_limit = 1000  # Límite de entradas en el caché

# Cargar caché desde archivo si existe, sino crear uno nuevo con límite
if os.path.exists(cache_file):
    with open(cache_file, 'rb') as f:
        response_cache = pickle.load(f)
    if not isinstance(response_cache, OrderedDict):
        response_cache = OrderedDict(response_cache)
else:
    response_cache = OrderedDict()

# Conectar a MongoDB
client = MongoClient('localhost', 27017)
db = client['email_database_metis2']
emails_collection = db['emails']

# Cargar modelo de embeddings
embedding_model = SentenceTransformer('paraphrase-multilingual-MiniLM-L12-v2')

# Crear índices solo si no existen
def initialize_collection():
    existing_indexes = {index['name'] for index in emails_collection.list_indexes()}
    if 'text_index' not in existing_indexes:
        emails_collection.create_index([
            ('from', TEXT),
            ('to', TEXT),
            ('subject', TEXT),
            ('body', TEXT),
            ('headers_text', TEXT),
            ('attachments', TEXT),
            ('attachments_content', TEXT),
            ('summary', TEXT),
            ('relevant_terms', TEXT),
            ('semantic_domain', TEXT)
        ], name='text_index', default_language='spanish')
    if 'date_index' not in existing_indexes:
        emails_collection.create_index([('date', ASCENDING)], name='date_index')
    if 'message_id_1' not in existing_indexes:
        emails_collection.create_index([('message_id', ASCENDING)], unique=True, name='message_id_1')
    if 'common_filters_index' not in existing_indexes:
        emails_collection.create_index([
            ('from', ASCENDING),
            ('to', ASCENDING),
            ('date', ASCENDING),
            ('semantic_domain', ASCENDING)
        ], name='common_filters_index')
    print("Índices creados o verificados en MongoDB.")

initialize_collection()

# Funciones de optimización de texto
def text_optimization(text):
    text = text.lower()
    text = re.sub('<[^<]+?>', '', text)
    text = re.sub(r'http\S+', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text[:10000]

def text_optimization_attachments(text):
    text = text.lower()
    text = re.sub('<[^<]+?>', '', text)
    text = re.sub(r'http\S+', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text[:1000]

def generate_embedding(subject, body, summary):
    """Genera un embedding comprimido para un correo electrónico."""
    text = f"{subject or ''} {body or ''} {summary or ''}".strip()
    if not text:
        return None
    try:
        embedding = embedding_model.encode(text).tolist()
        compressed_embedding = zlib.compress(np.array(embedding, dtype=np.float32).tobytes())
        return compressed_embedding
    except Exception as e:
        print(f"Error al generar embedding: {e}")
        return None

def infer_semantic_domain(subject, body, attachments_content):
    """Infiere el dominio semántico del correo usando Mistral."""
    text = f"{subject or ''} {body or ''} {' '.join(attachments_content) or ''}".strip()
    if not text:
        return 'general', 0.5

    prompt = f"""
    Analiza el siguiente texto de un correo electrónico y determina el dominio semántico más adecuado (ejemplo: "viajes", "negocios", "personal", "promociones", "técnico", "general").
    Devuelve EXCLUSIVAMENTE un objeto JSON con:
    - "semantic_domain": El dominio semántico identificado (cadena vacía si no se puede determinar).
    - "confidence": Valor entre 0 y 1 que indica la confianza en la selección.

    **Instrucciones**:
    1. Considera el contexto, palabras clave y propósito implícito del correo.
    2. Si no hay un dominio claro, devuelve "general" con confianza baja.
    3. Ejemplo: Para un correo sobre reservas de vuelos, el dominio sería "viajes".
    4. Devuelve solo un objeto JSON válido, sin comentarios ni texto adicional.

    **Ejemplo**:
    {{
        "semantic_domain": "viajes",
        "confidence": 0.95
    }}

    Texto:
    {text[:2000]}
    """

    response = call_mistral_api(prompt)
    try:
        cleaned_response = response.strip()
        if cleaned_response.startswith('```json'):
            cleaned_response = cleaned_response[7:].strip()
        if cleaned_response.endswith('```'):
            cleaned_response = cleaned_response[:-3].strip()
        result = json.loads(cleaned_response)
        return result.get('semantic_domain', 'general'), result.get('confidence', 0.5)
    except Exception as e:
        print(f"Error al inferir dominio semántico: {e}")
        return 'general', 0.5

def process_email_with_mistral(email):
    subject = email.get('subject', '')
    from_ = email.get('from', '')
    to = email.get('to', '')
    body = email.get('body', '')
    attachments = ' '.join([a for a in email.get('attachments', []) if a is not None])
    attachments_content = ' '.join([c for c in email.get('attachments_content', []) if c is not None])

    optimized_subject = text_optimization(subject)
    optimized_from = text_optimization(from_)
    optimized_to = text_optimization(to)
    optimized_body = text_optimization(body)
    optimized_attachments = text_optimization(attachments)
    optimized_attachments_content = text_optimization_attachments(attachments_content)

    text = f"{optimized_subject} {optimized_from} {optimized_to} {optimized_body} {optimized_attachments} {optimized_attachments_content}"

    prompt = f"""
    Analiza el siguiente texto de un correo electrónico y devuelve EXCLUSIVAMENTE un objeto JSON válido con:
    1. "summary": Un resumen con temas clave, nombres propios y términos específicos del dominio.
    2. "relevant_terms": Un diccionario donde las claves son términos relevantes y los valores son objetos con:
       - "frequency": Número entero de veces que aparece.
       - "context": Breve descripción de su significado o uso.
       - "type": "acción", "nombre_propio", "url" o "definición_temporal".
    
    NO INCLUYAS NINGÚN TEXTO FUERA DEL JSON, SOLO EL OBJETO JSON. SI NO PUEDES GENERAR UN JSON VÁLIDO, DEVUELVE {{"summary": "Error en el análisis", "relevant_terms": {{}}}}.
    Ejemplo:
    {{
        "summary": "Resumen del correo...",
        "relevant_terms": {{
            "reunión": {{"frequency": 1, "context": "Evento programado", "type": "acción"}},
            "Juan Pérez": {{"frequency": 1, "context": "Persona mencionada", "type": "nombre_propio"}}
        }}
    }}

    Texto:
    {text}
    """

    response = call_mistral_api(prompt)
    try:
        result = json.loads(response)
        if "relevant_terms" in result:
            repaired_terms = {}
            for term, data in result["relevant_terms"].items():
                if isinstance(data, dict) and "frequency" in data and "context" in data and "type" in data:
                    repaired_terms[term] = data
                else:
                    print(f"Ignorando término inválido: {term} -> {data}")
            result["relevant_terms"] = repaired_terms
        return result.get("summary", ""), result.get("relevant_terms", {})
    except json.JSONDecodeError:
        print(f"Error parsing JSON, intentando reparar: {response}")
        try:
            repaired_json = {"summary": "", "relevant_terms": {}}
            summary_match = re.search(r'(?:Resumen|Temas clave|El correo trata sobre|This appears to be)[\s\S]*?(?=\n\n|\Z)', response, re.IGNORECASE)
            if summary_match and "Error en el análisis" not in summary_match.group(0):
                repaired_json["summary"] = summary_match.group(0).strip()
            else:
                repaired_json["summary"] = "No se pudo generar un resumen válido" if "Error en el análisis" in response else response[:200].strip()
            
            terms = {}
            name_matches = re.findall(r'\b[A-Z][a-z]+\s[A-Z][a-z]+\b', response)
            for name in set(name_matches):
                terms[name] = {
                    "frequency": response.lower().count(name.lower()),
                    "context": "Persona mencionada en el correo",
                    "type": "nombre_propio"
                }
            words = re.findall(r'\b\w+\b', response.lower())
            word_counts = Counter(words)
            for word, count in word_counts.items():
                if count > 1 and len(word) > 3 and word not in terms:
                    terms[word] = {
                        "frequency": count,
                        "context": "Término frecuente en el correo",
                        "type": "acción" if word.endswith(('ar', 'er', 'ir')) else "definición_temporal"
                    }
            repaired_json["relevant_terms"] = terms
            
            if repaired_json["summary"] == "No se pudo generar un resumen válido" and not repaired_json["relevant_terms"]:
                return "No se pudo generar un resumen válido", {}
            return repaired_json["summary"], repaired_json["relevant_terms"]
        except Exception as e:
            print(f"No se pudo reparar el JSON: {e}")
            return "No se pudo generar un resumen válido", {}

def call_mistral_api(prompt):
    prompt_hash = hashlib.md5(prompt.encode('utf-8')).hexdigest()
    if prompt_hash in response_cache:
        print(f"Usando respuesta del caché para prompt hash: {prompt_hash}")
        response_cache.move_to_end(prompt_hash)
        return response_cache[prompt_hash]

    url = "http://localhost:11434/api/generate"
    payload = {
        "model": "mistral-custom",
        "prompt": prompt,
        "stream": False,
        "temperature": 0.7,
        "num_predict": 512,
        "num_ctx": 32768
    }
    
    try:
        response = requests.post(url, json=payload)
        response.raise_for_status()
        result = response.json()['response']
        if len(response_cache) >= cache_limit:
            response_cache.popitem(last=False)
        response_cache[prompt_hash] = result
        with open(cache_file, 'wb') as f:
            pickle.dump(response_cache, f)
        return result
    except requests.RequestException as e:
        print(f"Error al contactar con la API: {e}")
        return f"Error al contactar con la API: {e}"
    except json.JSONDecodeError:
        print("Error al decodificar la respuesta JSON")
        return "Error al decodificar la respuesta JSON"

# Autenticación con Gmail
if os.path.exists(TOKEN_FILE):
    with open(TOKEN_FILE, 'rb') as token:
        creds = pickle.load(token)
else:
    flow = InstalledAppFlow.from_client_secrets_file(CREDS_FILE, SCOPES, redirect_uri='http://localhost:8080')
    auth_url, _ = flow.authorization_url(prompt='consent')
    print("Please go to this URL and authorize access:", auth_url)
    code = input('Enter the authorization code: ')
    flow.fetch_token(code=code)
    creds = flow.credentials
    with open(TOKEN_FILE, 'wb') as token:
        pickle.dump(creds, token)

service = build('gmail', 'v1', credentials=creds)

def download_attachment(service, user_id, msg_id, part):
    attachment = service.users().messages().attachments().get(userId=user_id, messageId=msg_id, id=part['body']['attachmentId']).execute()
    file_data = base64.urlsafe_b64decode(attachment['data'].encode('UTF-8'))
    filename = part.get('filename', 'unnamed_attachment')
    path = f'./attachments/{filename}'
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'wb') as f:
        f.write(file_data)
    return path

def analyze_attachment(file_path):
    mime = mimetypes.guess_type(file_path)[0]
    if mime is None:
        return f"No se pudo determinar el tipo MIME del archivo: {file_path}"
    elif mime == 'text/plain':
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            return f"Error leyendo archivo de texto: {str(e)}"
    elif mime == 'application/pdf':
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
            return text
        except FileNotDecryptedError:
            return "El archivo PDF está encriptado y no pudo ser analizado."
        except Exception as e:
            return f"Error al analizar PDF: {str(e)}"
    elif mime.startswith('image/'):
        try:
            text = pytesseract.image_to_string(Image.open(file_path))
            return text
        except Exception as e:
            return f"OCR error: {str(e)}"
    elif mime == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
        try:
            doc = DocxDocument(file_path)
            return '\n'.join([paragraph.text for paragraph in doc.paragraphs])
        except Exception as e:
            return f"Error al analizar documento Word: {str(e)}"
    elif mime == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            text = ""
            for row in sheet.iter_rows(values_only=True):
                text += ', '.join([str(cell) for cell in row]) + '\n'
            return text
        except Exception as e:
            return f"Error al analizar documento Excel: {str(e)}"
    else:
        return f"Tipo de archivo no soportado: {mime}"

def extract_urls(text):
    url_pattern = re.compile(r'https?://\S+')
    urls = url_pattern.findall(text)
    return [{'action': 'visit', 'description': 'Link encontrado', 'url': url} for url in urls]

# Procesar e insertar correos uno por uno
def process_and_insert_email(msg, service, user_id):
    headers = {header['name']: header['value'] for header in msg['payload']['headers']}
    subject = headers.get('Subject', 'No Subject')
    from_ = headers.get('From', 'Unknown')
    to = headers.get('To', 'Unknown')
    date = headers.get('Date', 'Unknown')
    message_id = headers.get('Message-ID', 'Unknown')
    
    if message_id == 'Unknown':
        hash_string = f"{subject}{from_}{to}{date}".encode('utf-8')
        message_id = hashlib.sha256(hash_string).hexdigest()
    
    # Verificar si el correo ya existe
    if emails_collection.find_one({'message_id': message_id}):
        print(f"Correo con message_id {message_id} ya existe, omitiendo.")
        return
    
    print(f"Procesando correo: From: {from_}, To: {to}, Subject: {subject}, Date: {date}")

    body = ''
    headers_text = ''
    if 'parts' in msg['payload']:
        for part in msg['payload']['parts']:
            if part['mimeType'] == 'text/plain':
                data = part['body'].get('data')
                if data:
                    body = base64.urlsafe_b64decode(data.encode('ASCII')).decode('utf-8')
                    break
            headers_text += f"{part['mimeType']}\n"
    else:
        data = msg['payload']['body'].get('data')
        if data:
            body = base64.urlsafe_b64decode(data.encode('ASCII')).decode('utf-8')
        headers_text = '\n'.join([f"{h['name']}: {h['value']}" for h in msg['payload']['headers']])

    attachments = []
    attachments_content = []
    if 'parts' in msg['payload']:
        for part in msg['payload']['parts']:
            if part.get('filename'):
                attachments.append(part['filename'])
                if part.get('body', {}).get('attachmentId'):
                    attachment_path = download_attachment(service, user_id, msg['id'], part)
                    analyzed_content = analyze_attachment(attachment_path)
                    attachments_content.append(analyzed_content)

    full_text = f"{headers_text}\n{subject}\n{from_}\n{to}\n{body}\n{' '.join(attachments_content)}"
    
    email_dict = {
        'subject': subject,
        'from': from_,
        'to': to,
        'body': body,
        'attachments': attachments,
        'attachments_content': attachments_content
    }
    
    summary, relevant_terms = process_email_with_mistral(email_dict)
    semantic_domain, domain_confidence = infer_semantic_domain(subject, body, attachments_content)

    # Verificar si el JSON es válido antes de continuar
    if summary == "No se pudo generar un resumen válido" and not relevant_terms:
        print(f"JSON no válido para message_id {message_id}, omitiendo inserción en MongoDB.")
        return

    # Generar embedding
    embedding = generate_embedding(subject, body, summary)
    
    urls = extract_urls(full_text)
    
    email_document = {
        'from': from_,
        'to': to,
        'subject': subject,
        'date': date,
        'body': body,
        'headers_text': headers_text,
        'attachments': attachments,
        'attachments_content': attachments_content,
        'message_id': message_id,
        'urls': urls,
        'summary': summary,
        'relevant_terms': relevant_terms,
        'embedding': embedding,
        'semantic_domain': semantic_domain,
        'domain_confidence': domain_confidence
    }
    
    # Insertar el correo inmediatamente en MongoDB
    try:
        emails_collection.insert_one(email_document)
        print(f"Insertado correo con message_id: {message_id}")
    except Exception as e:
        print(f"Error al insertar correo con message_id {message_id}: {e}")

# Descargar y procesar correos
labels = ['INBOX', 'SENT']
for label in labels:
    results = service.users().messages().list(userId='me', labelIds=[label], maxResults=250).execute()
    messages = results.get('messages', [])
    
    for i, message in enumerate(messages, 1):
        print(f"Descargando mensaje {i}/{len(messages)} de {label}...")
        msg = service.users().messages().get(userId='me', id=message['id']).execute()
        process_and_insert_email(msg, service, 'me')

print("Procesamiento de correos completado.")
